"""Build the one-page Arabic+English instruction sheet for the blind
re-annotation task (`blind_reannotation_sample.xlsx`).

Uses the project's shared RTL helpers (CLAUDE.md: every .docx builder here
must use experiments/pilot_3models/rtl_utils.py, never hand-rolled bidi).
Each paragraph's direction is decided independently via
apply_paragraph_direction() - Arabic paragraphs render RTL, the English
paragraphs stay LTR, matching the file's actual bilingual structure.
"""
import sys
from pathlib import Path

PILOT_DIR = Path(__file__).resolve().parents[4] / "experiments" / "pilot_3models"
if str(PILOT_DIR) not in sys.path:
    sys.path.insert(0, str(PILOT_DIR))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from rtl_utils import apply_paragraph_direction

from common import IAA_DIR

AR_TITLE = "ورقة تعليمات: مهمة التصنيف المستقل (للتحقق من ثبات التصنيف)"
EN_TITLE = "Instruction Sheet: Independent Annotation Task (Agreement / Reliability Check)"

AR_BODY = [
    "الهدف: تعبئة ورقة العمل المرفقة (blind_reannotation_sample.xlsx) بشكل مستقل تمامًا، "
    "دون الرجوع إلى أي تصنيف سابق لنفس الأبيات (سواء كان تصنيفك الأصلي أو نتائج النماذج). "
    "هذه الورقة لا تحتوي على أي تسمية مسبقة، فقط نص البيت.",
    "الخطوات: لكل بيت في تبويب \"Annotation\"، حدد لكل عنصر من العناصر الأربعة "
    "(مجاز / تشبيه / كناية (periphrasis) / طباق (antithesis)) ما إذا كان موجودًا "
    "(نعم / لا)، وإن كان موجودًا، اكتب الموضع (التعبير) الدقيق من البيت الذي يحتوي عليه. "
    "يمكن أن يحتوي بيت واحد على أكثر من عنصر، أو أكثر من موضع واحد لنفس العنصر - "
    "في هذه الحالة استخدم خلية الملاحظات لتوضيح ذلك.",
    "مرجع التعريفات: تبويب \"Element Definitions\" في نفس الملف يحتوي على التعريفات "
    "والأمثلة الأصلية (نفس التعريفات المستخدمة في بناء هذا المشروع) - الرجاء الرجوع إليها "
    "عند الشك، لا الاعتماد على الذاكرة فقط.",
    "مهم: إذا كنت تُعيد تصنيف أبيات قمت بتصنيفها من قبل، الرجاء عدم فتح أو مراجعة ملف "
    "التصنيف الأصلي أثناء العمل على هذه الورقة - الفكرة هي قياس مدى الاتساق عند إعادة "
    "التصنيف بعد فترة زمنية، دون الاعتماد على تذكر القرار السابق.",
    "الوقت المتوقع: 80 بيتًا، نحو 1-2 دقيقة للبيت الواحد في المتوسط (أكثر للأبيات التي "
    "تحتوي على أكثر من عنصر) - أي ما يقارب 1.5-2.5 ساعة إجمالاً، يمكن تقسيمها على عدة جلسات.",
]

EN_BODY = [
    "Goal: fill in the attached worksheet (blind_reannotation_sample.xlsx) completely "
    "independently, without referring back to any prior labeling of the same verses - "
    "neither your own original annotation nor any model output. This sheet carries no "
    "pre-filled label of any kind, only the verse text.",
    "Steps: for every verse in the \"Annotation\" tab, decide for each of the four "
    "elements (Trope / Simile / Kenning=Periphrasis / Antonymy=Antithesis) whether it is "
    "present (Yes/No), and if present, write the exact span (expression) from the verse "
    "that contains it. A single verse may contain more than one element, or more than "
    "one instance of the same element - use the Notes cell to note that.",
    "Definitions reference: the \"Element Definitions\" tab in the same file holds the "
    "original definitions and worked examples (the same ones used throughout this "
    "project) - please consult them when in doubt rather than relying on memory alone.",
    "Important: if you are re-annotating verses you already annotated before, please do "
    "not open or review the original annotation file while working on this sheet - the "
    "point is to measure how consistent a re-annotation is after a time gap, not how well "
    "you can recall the earlier decision.",
    "Expected time: 80 verses, roughly 1-2 minutes per verse on average (more for verses "
    "with multiple elements) - about 1.5-2.5 hours total, can be split across sessions.",
]


def add_heading(doc, text_ar, text_en):
    p = doc.add_paragraph()
    run = p.add_run(text_ar)
    run.bold = True
    run.font.size = Pt(15)
    apply_paragraph_direction(p, text_ar)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p2 = doc.add_paragraph()
    run2 = p2.add_run(text_en)
    run2.bold = True
    run2.font.size = Pt(13)
    apply_paragraph_direction(p2, text_en)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_body_block(doc, items, lang_label):
    label_p = doc.add_paragraph()
    r = label_p.add_run(lang_label)
    r.bold = True
    r.font.size = Pt(12)
    apply_paragraph_direction(label_p, lang_label)

    for i, text in enumerate(items, start=1):
        p = doc.add_paragraph()
        run = p.add_run(f"{i}. {text}")
        run.font.size = Pt(10.5)
        apply_paragraph_direction(p, text)
        p.paragraph_format.space_after = Pt(6)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin = section.right_margin = Cm(1.8)

    add_heading(doc, AR_TITLE, EN_TITLE)
    doc.add_paragraph()
    add_body_block(doc, AR_BODY, "بالعربية:")
    doc.add_paragraph()
    add_body_block(doc, EN_BODY, "In English:")

    out_path = IAA_DIR / "sheets" / "instructions.docx"
    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
