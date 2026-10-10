"""Build paper_draft_for_supervisors.docx from paper_trimmed.md.

Reuses the project's shared RTL helpers (experiments/pilot_3models/rtl_utils.py)
so Arabic paragraphs/runs get <w:bidi/>/<w:rtl/> exactly as every other
.docx builder in this project does. Markdown-to-docx conversion follows the
same line-based approach as experiments/pilot_3models/build_report.py.
"""
import sys
import re
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "experiments" / "pilot_3models"))

import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from rtl_utils import apply_paragraph_direction

SRC = Path(__file__).resolve().parent / "paper_trimmed.md"
OUT = Path(__file__).resolve().parent / "paper_draft_for_supervisors.docx"


def _is_separator_row(cells):
    return all(re.fullmatch(r":?-+:?", c.strip()) for c in cells if c.strip() != "")


def _parse_table_block(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not _is_separator_row(cells):
            rows.append(cells)
        i += 1
    return rows, i


def _render_inline(paragraph, text):
    """Render **bold**, *italic*, and `code` spans as real runs, applying
    RTL direction to the whole paragraph based on its own text."""
    # tokenize on **bold**, *italic*, `code` without nesting
    tokens = re.split(r'(\*\*.*?\*\*|\*[^*]+\*|`[^`]+`)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("*") and tok.endswith("*") and not tok.startswith("**"):
            run = paragraph.add_run(tok[1:-1])
            run.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(tok)
    apply_paragraph_direction(paragraph, text)


def build():
    md = SRC.read_text(encoding="utf-8")
    lines = md.splitlines()

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    i = 0
    n = len(lines)
    title_done = False
    while i < n:
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            # Document title
            heading_text = line[2:].strip()
            p = doc.add_heading(level=0)
            _render_inline(p, heading_text)
            title_done = True
            i += 1
            continue

        if line.startswith("**Author(s):**"):
            p = doc.add_paragraph()
            _render_inline(p, line)
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_heading(level=2)
            _render_inline(p, line[4:].strip())
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_heading(level=1)
            _render_inline(p, line[3:].strip())
            i += 1
            continue

        if line.strip().startswith("|"):
            rows, next_i = _parse_table_block(lines, i)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for r, row_vals in enumerate(rows):
                    for c, val in enumerate(row_vals):
                        cell_text = val.replace("**", "")
                        cell = table.cell(r, c)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        _render_inline(p, cell_text)
                        if r == 0:
                            for run in p.runs:
                                run.bold = True
            i = next_i
            doc.add_paragraph()  # spacing after table
            continue

        if line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _render_inline(p, line.strip()[2:])
            i += 1
            continue

        if re.match(r"^\*.*\*$", line.strip()) and line.strip().count("*") == 2:
            # a whole-line italic note, e.g. *Source: ...*
            p = doc.add_paragraph()
            _render_inline(p, line.strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        # plain paragraph
        p = doc.add_paragraph()
        _render_inline(p, line.strip())
        i += 1

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
