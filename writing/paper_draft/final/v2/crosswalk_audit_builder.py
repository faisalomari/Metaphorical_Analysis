"""Builds crosswalk_audit_sheet.xlsx: a blind spot-check sample of the
53-category classical-rhetoric-taxonomy -> 4-element crosswalk that
experiments/pilot_3models/ground_truth.py applies to every poem's ground
truth (SUBSTRING_MAP: "مجاز"/"استعارة" -> Trope, "تشبيه" -> Simile,
"كناية" -> Kenning/Periphrasis, "طباق" -> Antonymy/Antithesis; anything
else -> not evidence for any of the 4 elements).

Answers reviewer_simulation.md Reviewer 1 Sec.1 / ranked-fix #8: "spot-check
the 53-to-4-category crosswalk against Prof. Hussein's own top-level
judgment." This builds the blind instrument only -- it does not and cannot
perform the check itself, since that requires Prof. Hussein's own judgment.

Zero API calls. Reads only the five original ground-truth spreadsheets
(data/ground_truth/ali_ana/original/*.xlsx), unmodified, to enumerate every
distinct (primary-label, secondary-label) pair actually used across all 5
poems, and applies the project's own mapping rule (copied here verbatim
from ground_truth.py's SUBSTRING_MAP, not re-derived or altered).

NOTE on scope: this project's prose describes the taxonomy as having
"fifty-three" categories (paper_trimmed.md, Data and Methods). Direct
enumeration here finds 59 distinct (primary, secondary) label pairs across
the 5 poems' annotation sheets -- see leakage_quantification.md / this
task's final report for this discrepancy, which is flagged, not resolved,
by this script.
"""
import random
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[4]
GT_DIR = ROOT / "data" / "ground_truth" / "ali_ana" / "original"

POEM_FILES = {
    "Imru1": "امرؤ القيس 1.xlsx",
    "Imru2": "امرؤ القيس 8.xlsx",
    "Nabigha": "النابغة الذبياني.xlsx",
    "Jarir": "جرير.xlsx",
    "Abu_Tammam": "ابو تمام.xlsx",
}

# Copied verbatim from experiments/pilot_3models/ground_truth.py -- not
# altered. See that file's own docstring for the mapping rationale.
SUBSTRING_MAP = {
    "Trope (مجاز)": ["مجاز", "استعارة"],
    "Simile (تشبيه)": ["تشبيه"],
    "Periphrasis (كناية)": ["كناية"],
    "Antithesis (طباق)": ["طباق"],
}

OUT_PATH = Path(__file__).resolve().parent / "crosswalk_audit_sheet.xlsx"
SAMPLE_N = 40
SEED = 20261007
FORCE_INCLUDE = [("استعارة", "استعارة تمثيلية مبنية على الكناية")]  # dual-mapped edge case


def map_label(primary: str, secondary: str):
    labels = f"{primary} {secondary}"
    return [name for name, subs in SUBSTRING_MAP.items() if any(s in labels for s in subs)]


def collect_unique_pairs():
    seen = {}
    for poem, fname in POEM_FILES.items():
        df = pd.read_excel(GT_DIR / fname)
        for _, row in df.iterrows():
            primary = str(row.get("العنصر البلاغي المستعمل", "") or "")
            secondary = str(row.get("العنصر البلاغي الثانوي", "") or "")
            if primary == "nan":
                primary = ""
            if secondary == "nan":
                secondary = ""
            if not primary and not secondary:
                continue
            key = (primary, secondary)
            seen[key] = seen.get(key, 0) + 1
    return seen  # {(primary, secondary): count_across_all_5_poems}


def build_sample(pairs: dict):
    keys = sorted(pairs.keys())
    rng = random.Random(SEED)
    sample = rng.sample(keys, min(SAMPLE_N, len(keys)))
    for forced in FORCE_INCLUDE:
        if forced not in sample and forced in keys:
            sample[rng.randrange(len(sample))] = forced
    rng.shuffle(sample)  # blind order -- not sorted by frequency or by element
    return sample


def mapped_display(primary, secondary):
    mapped = map_label(primary, secondary)
    if not mapped:
        return "— (not mapped to any of the 4 elements)"
    return " + ".join(mapped)


def write_instructions(wb):
    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_view.rightToLeft = True
    bold = Font(bold=True, size=13)
    normal = Font(size=11)
    ar_font = Font(size=12, name="Calibri")

    lines_ar = [
        ("تعليمات — فحص عشوائي لمطابقة فئات البلاغة مع العناصر الأربعة", bold),
        ("", normal),
        ("هذا المشروع يرسم أكثر من خمسين فئة بلاغية فرعية (كما وردت في جداول تحليلك)", ar_font),
        ("إلى أربعة عناصر رئيسية فقط: مجاز (Trope)، تشبيه (Simile)، كناية (Periphrasis)، طباق (Antithesis).", ar_font),
        ("القاعدة المستخدمة آلية بالكامل: أي تسمية تحتوي على \"مجاز\" أو \"استعارة\" تُحسب كـ Trope،", ar_font),
        ("وأي تسمية تحتوي على \"تشبيه\" تُحسب كـ Simile، وهكذا — دون أي تدخل بشري إضافي.", ar_font),
        ("", normal),
        ("المطلوب منك: في ورقة \"Audit Sample\"، لكل صف، تجد فقط تسمية الفئة (كما كتبتها أنت في تحليلك)", ar_font),
        ("والعنصر الذي حُسبت له آليًا. لا تظهر لك أي أرقام أبيات أو تكرارات أو اسم القصيدة — هذا فحص", ar_font),
        ("\"أعمى\" مقصود، لتحكم على التسمية بذاتها لا على مدى تكرارها.", ar_font),
        ("", normal),
        ("ضع علامة ✓ في خانة \"Agree\" إذا كان التصنيف الآلي صحيحًا من منظورك كخبير،", ar_font),
        ("أو علامة ✓ في خانة \"Disagree\" إذا كنت ترى أن التصنيف خاطئ أو ناقص،", ar_font),
        ("واكتب أي ملاحظة في خانة \"Notes\" — خصوصًا إن كانت التسمية غامضة أو تنتمي لعنصرين معًا.", ar_font),
        ("", normal),
        ("ملاحظة خاصة: الصف المعلّم بعلامة * يمثّل حالة تُحسب حاليًا لعنصرين في آنٍ واحد", ar_font),
        ("(مجاز + كناية) — نرجو تعليقك على هذه الحالة بالتحديد.", ar_font),
        ("", normal),
        ("English instructions", bold),
        ("", normal),
        ("This project maps 50+ fine-grained rhetorical-category labels (as you recorded them", normal),
        ("in your own annotation sheets) down to just four top-level elements: Trope (مجاز),", normal),
        ("Simile (تشبيه), Periphrasis (كناية), Antithesis (طباق). The mapping rule is purely", normal),
        ("mechanical: any label containing the substring \"مجاز\" or \"استعارة\" is counted as", normal),
        ("Trope, any label containing \"تشبيه\" as Simile, and so on -- with no further human", normal),
        ("judgment applied at this step.", normal),
        ("", normal),
        ("On the \"Audit Sample\" tab, each row shows only the category label (exactly as you", normal),
        ("wrote it) and the element it was automatically mapped to. No verse numbers, counts,", normal),
        ("or poem names are shown -- this is a deliberately blind check, so the judgment rests", normal),
        ("on the label alone, not on how often it occurs.", normal),
        ("", normal),
        ("Please mark \"Agree\" if the automatic mapping matches your own expert judgment, or", normal),
        ("\"Disagree\" if you believe it is wrong or incomplete, and use \"Notes\" for anything", normal),
        ("ambiguous -- especially a label you feel genuinely belongs to more than one element.", normal),
        ("", normal),
        ("One row is marked with * -- this label is currently counted toward TWO elements at", normal),
        ("once (Trope + Periphrasis). Your judgment on this specific case would be especially", normal),
        ("useful.", normal),
    ]
    for i, (text, font) in enumerate(lines_ar, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, horizontal="right" if any("؀" <= c <= "ۿ" for c in text) else "left")
    ws.column_dimensions["A"].width = 110


def write_sample_sheet(wb, sample, pairs):
    ws = wb.create_sheet("Audit Sample")
    ws.sheet_view.rightToLeft = True
    headers = ["#", "Category label (التسمية)", "Mapped element (العنصر المحسوب)", "Agree", "Disagree", "Notes"]
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    for i, (primary, secondary) in enumerate(sample, start=1):
        label = secondary if secondary and secondary != primary else primary
        if secondary and primary and secondary != primary:
            label = f"{secondary}  (أصل: {primary})"
        is_forced = (primary, secondary) in FORCE_INCLUDE
        row_num_display = f"{i}*" if is_forced else str(i)
        ws.cell(row=i + 1, column=1, value=row_num_display)
        ws.cell(row=i + 1, column=2, value=label).alignment = Alignment(horizontal="right", wrap_text=True)
        ws.cell(row=i + 1, column=3, value=mapped_display(primary, secondary)).alignment = Alignment(horizontal="center", wrap_text=True)
        # Agree/Disagree/Notes left empty for Prof. Hussein to fill in.

    widths = [5, 48, 32, 10, 10, 40]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"


def main():
    pairs = collect_unique_pairs()
    print(f"Total distinct (primary, secondary) category pairs found: {len(pairs)}")
    sample = build_sample(pairs)
    print(f"Sample size: {len(sample)}")

    wb = Workbook()
    write_instructions(wb)
    write_sample_sheet(wb, sample, pairs)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
