"""Task 1 (2026-07-26 session) - per-verse chat-like reasoning document.

Assembles Config C's EXISTING stage-1 Fable 5 free-form Arabic rhetorical
analyses (see PROJECT_STATE.md Phase 2 SS3i / results/phase2_chatlike/report_3i.md)
into a single readable .docx for the supervisors: verse text, the model's full
analysis verbatim (both independent runs), then Dr. Ali Hussein's ground-truth
annotation for comparison.

Zero new API calls - purely a presentation layer over already-collected JSON.
Covers:
  - Imru2 (poem 8), all 23 verses, Config C Stage 1, both runs
  - The 11 hand-picked Nabigha verses from the targeted مجاز مرسل recall probe
    (Phase 2 SS3i addendum / report_3i.md SS8), both runs

Uses experiments/pilot_3models/rtl_utils.py for all RTL/bidi handling per
CLAUDE.md's standing rule - no hand-rolled bidi logic here.
"""
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "experiments" / "pilot_3models"))

from docx import Document
from docx.shared import Pt

from rtl_utils import apply_paragraph_direction, is_majority_arabic
from ground_truth import load_ground_truth_spans

RESULTS = ROOT / "results" / "phase2_chatlike"
OUT_PATH = RESULTS / "verse_by_verse_reasoning.docx"

ELEMENT_LABEL = {
    "Trope": "Trope (مجاز)",
    "Simile": "Simile (تشبيه)",
    "Kenning": "Periphrasis (كناية)",
    "Antonymy": "Antithesis (طباق)",
}
ELEMENT_ORDER = ["Trope", "Simile", "Kenning", "Antonymy"]

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def add_markdown_block(doc, text):
    """Render a Fable 5 stage-1 analysis (light markdown: #/##/### headers,
    **bold**, --- rules) as docx paragraphs, applying RTL per-paragraph."""
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            doc.add_paragraph()
            continue
        if set(line) == {"-"}:
            continue  # horizontal rule marker, not meaningful in docx
        heading_match = re.match(r"^(#{1,4})\s*(.*)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            content = heading_match.group(1) and heading_match.group(2)
            content = BOLD_RE.sub(r"\1", content).strip()
            p = doc.add_heading(level=min(level + 2, 6))
            p.add_run(content)
            apply_paragraph_direction(p, content)
            continue
        p = doc.add_paragraph()
        pos = 0
        for m in BOLD_RE.finditer(line):
            if m.start() > pos:
                p.add_run(line[pos:m.start()])
            b = p.add_run(m.group(1))
            b.bold = True
            pos = m.end()
        if pos < len(line):
            p.add_run(line[pos:])
        apply_paragraph_direction(p, line)


def add_verse_text(doc, verse_text):
    p = doc.add_paragraph()
    r = p.add_run(verse_text.strip())
    r.bold = True
    r.font.size = Pt(13)
    apply_paragraph_direction(p, verse_text)
    p.paragraph_format.space_after = Pt(6)


def add_ground_truth(doc, gt_for_verse):
    label_p = doc.add_paragraph()
    lr = label_p.add_run("Ground truth (Dr. Ali Hussein):")
    lr.italic = True
    apply_paragraph_direction(label_p, "Ground truth (Dr. Ali Hussein):")

    any_gt = False
    for el in ELEMENT_ORDER:
        instances = gt_for_verse.get(el, [])
        if not instances:
            continue
        any_gt = True
        spans = "; ".join(
            f"«{inst['span']}»" + (f" ({inst['secondary']})" if inst.get("secondary") else "")
            for inst in instances
        )
        line = f"{ELEMENT_LABEL[el]}: {spans}"
        p = doc.add_paragraph()
        p.add_run(line)
        apply_paragraph_direction(p, line)
    if not any_gt:
        p = doc.add_paragraph()
        p.add_run("(no ground-truth rhetorical instance annotated for this verse)")
        apply_paragraph_direction(p, "none")


def build():
    gt_imru2 = load_ground_truth_spans("Imru2")
    gt_nabigha = load_ground_truth_spans("Nabigha")

    run1 = {x["verse_no"]: x for x in json.loads(
        (RESULTS / "config_c" / "Imru2" / "stage1_analysis_run1.json").read_text(encoding="utf-8"))}
    run2 = {x["verse_no"]: x for x in json.loads(
        (RESULTS / "config_c" / "Imru2" / "stage1_analysis_run2.json").read_text(encoding="utf-8"))}

    probe = json.loads((RESULTS / "config_c_probe" / "Nabigha" / "probe_results.json").read_text(encoding="utf-8"))

    doc = Document()

    title = doc.add_heading(level=0)
    title.add_run("Per-Verse Chat-Like Reasoning — Config C (Fable 5, Stage 1)")
    apply_paragraph_direction(title, "Per-Verse Chat-Like Reasoning")

    intro = (
        "Purpose: in response to the 2026-07-23 supervisor request to \"check line by line "
        "in chat and Claude Code, not using the API,\" this document assembles Config C's "
        "existing Stage-1 free-form Arabic rhetorical analyses (Fable 5, effort=medium, one "
        "verse per call, full poem as context) verbatim, verse by verse, alongside the verse "
        "text and Dr. Ali Hussein's ground-truth annotation. Every analysis below was already "
        "collected for Phase 2 SS3i (see PROJECT_STATE.md and results/phase2_chatlike/report_3i.md); "
        "this is a presentation layer only — zero new API calls were made to build it. "
        "Terminology follows GLOSSARY.md's 2026-07-23 corrections (antithesis, periphrasis; "
        "مجاز/trope as an umbrella over استعارة/تمثيل/مجاز مرسل subtypes)."
    )
    p = doc.add_paragraph(intro)
    p.paragraph_format.space_after = Pt(12)

    note1 = (
        "Section 1 (Imru2, poem 8): the full poem, all 23 verses, both Config C runs — "
        "this is the poem where Config C's Kenning/periphrasis recall gain was found "
        "(strict F1 0.111→0.316) and where it catches the motivating verse (verse 6) "
        "that every strict-JSON configuration (Existing, Config A, Config B) misses."
    )
    p = doc.add_paragraph(note1)
    p.italic = True

    doc.add_heading("Section 1 — Imru2 (Poem 8), all 23 verses", level=1)

    for no in range(1, 24):
        h = doc.add_heading(f"Verse {no}", level=2)
        apply_paragraph_direction(h, f"Verse {no}")
        add_verse_text(doc, run1[no]["verse_text"])
        add_ground_truth(doc, gt_imru2[no])

        r1_label = doc.add_paragraph()
        r1r = r1_label.add_run("Model reasoning — Run 1 (Fable 5, Stage 1, verbatim):")
        r1r.italic = True
        apply_paragraph_direction(r1_label, "Model reasoning Run 1")
        add_markdown_block(doc, run1[no]["analysis"])

        r2_label = doc.add_paragraph()
        r2r = r2_label.add_run("Model reasoning — Run 2 (Fable 5, Stage 1, verbatim):")
        r2r.italic = True
        apply_paragraph_direction(r2_label, "Model reasoning Run 2")
        add_markdown_block(doc, run2[no]["analysis"])

        doc.add_page_break()

    doc.add_heading("Section 2 — Nabigha, targeted مجاز مرسل recall probe (11 hand-picked verses)", level=1)
    note2 = (
        "These 11 verses are NOT a random or full-poem sample: 9 contain Nabigha's entire "
        "مجاز مرسل (metonymy) ground truth, plus 2 contrast verses for subtypes Phase 1 "
        "already found tractable (استعارة فعل, مجاز عقلي). See PROJECT_STATE.md Phase 2 SS3i "
        "addendum and report_3i.md SS8 for the full recall-probe methodology and results table "
        "— no precision/F1 claim is made from this probe, only per-instance hit/miss on "
        "known-positive verses."
    )
    p = doc.add_paragraph(note2)
    p.italic = True

    for no_str in sorted(probe.keys(), key=int):
        no = int(no_str)
        entries = probe[no_str]
        h = doc.add_heading(f"Verse {no} (Nabigha)", level=2)
        apply_paragraph_direction(h, f"Verse {no}")
        add_verse_text(doc, entries[0]["verse_text"])
        add_ground_truth(doc, gt_nabigha[no])

        for entry in sorted(entries, key=lambda e: e["run_index"]):
            label = doc.add_paragraph()
            lr = label.add_run(f"Model reasoning — Run {entry['run_index']} (Fable 5, Stage 1, verbatim):")
            lr.italic = True
            apply_paragraph_direction(label, "Model reasoning")
            add_markdown_block(doc, entry["stage1_analysis"])

        doc.add_page_break()

    doc.save(OUT_PATH)
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    build()
