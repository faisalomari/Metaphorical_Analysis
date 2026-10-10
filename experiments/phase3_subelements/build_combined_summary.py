"""Merges summary_dev.json + summary_test.json into one xlsx: per-subtype
confirmation rate across the full 551-row ground truth (DEV+TEST combined),
2 runs each. Run after both `run_subtype_check.py --split dev` and
`--split test` have completed.
"""
import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from phase3_config import RESULTS_DIR
from subtype_defs import SUBTYPE_DEFS

dev = json.load(open(RESULTS_DIR / "summary_dev.json", encoding="utf-8"))
test = json.load(open(RESULTS_DIR / "summary_test.json", encoding="utf-8"))
dev_map = {r["subtype"]: r for r in dev["subtypes"]}
test_map = {r["subtype"]: r for r in test["subtypes"]}
all_keys = list(dict.fromkeys(list(dev_map.keys()) + list(test_map.keys())))

rows = []
for k in all_keys:
    d = dev_map.get(k, {"n_rows": 0, "n_yes_run1": 0, "n_yes_run2": 0})
    t = test_map.get(k, {"n_rows": 0, "n_yes_run1": 0, "n_yes_run2": 0})
    n = d["n_rows"] + t["n_rows"]
    y1 = d["n_yes_run1"] + t["n_yes_run1"]
    y2 = d["n_yes_run2"] + t["n_yes_run2"]
    pct = round(100 * (y1 + y2) / (2 * n), 1) if n else 0
    parent = SUBTYPE_DEFS.get(k, {}).get("parent", "")
    rows.append([k, parent, n, y1, y2, pct])

rows.sort(key=lambda r: -r[2])

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "combined_summary"
ws.sheet_view.rightToLeft = True
ws.append(["العنصر_الفرعي", "العنصر_الرئيسي", "عدد_الأسطر", "مؤكد_تشغيلة1", "مؤكد_تشغيلة2", "نسبة_التأكيد_%"])
for r in rows:
    ws.append(r)
for i in range(1, 7):
    ws.column_dimensions[get_column_letter(i)].width = 26

total_cost = dev["cost_usd"] + test["cost_usd"]
total_calls = dev["total_calls"] + test["total_calls"]
ws2 = wb.create_sheet("run_info")
ws2.sheet_view.rightToLeft = True
ws2.append(["المقياس", "القيمة"])
ws2.append(["الموديل", dev["model"]])
ws2.append(["عدد الأسطر (DEV+TEST)", sum(r[2] for r in rows)])
ws2.append(["إجمالي الاستدعاءات", total_calls])
ws2.append(["التكلفة الإجمالية ($)", round(total_cost, 4)])
for i in range(1, 3):
    ws2.column_dimensions[get_column_letter(i)].width = 30

out_path = RESULTS_DIR / "combined_summary.xlsx"
wb.save(out_path)
print(f"written to {out_path}")
print(f"total rows={sum(r[2] for r in rows)}  total_calls={total_calls}  total_cost=${total_cost:.4f}")
