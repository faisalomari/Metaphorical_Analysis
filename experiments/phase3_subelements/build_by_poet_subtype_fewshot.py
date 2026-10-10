"""Same matrix as build_by_poet_subtype.py (poet x subtype, n/TP/FN/Recall),
but sourced from the few-shot run (raw/<subtype>_fewshot.jsonl) instead of
one-shot. Only covers the 13 subtypes eligible for few-shot (n>=10); each
subtype's n here is 2 less than the one-shot matrix's, since the 2 held-out
demonstration rows are excluded from testing. No new API calls.
"""
import json

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from phase3_config import RESULTS_DIR
from subtype_defs import SUBTYPE_DEFS
from build_by_poet import parent_short, majority, PARENT_ORDER
from run_fewshot_check import ELIGIBLE_SUBTYPES
from build_fewshot_comparison import slugify

RAW_DIR = RESULTS_DIR / "raw"


def main():
    agg = {}
    poets_seen = set()
    for subtype_key in ELIGIBLE_SUBTYPES:
        slug = slugify(subtype_key)
        path = RAW_DIR / f"{slug}_fewshot.jsonl"
        if not path.exists():
            print(f"  ! missing {path}, skipping {subtype_key}")
            continue
        by_row = {}
        with open(path, encoding="utf-8") as f:
            for line in f:
                rec = json.loads(line)
                key = (rec["poem_key"], rec["row_id"])
                by_row.setdefault(key, {"poet": rec["poet"]})[rec["run_idx"]] = rec.get("exists")
        per_poet = {}
        for key, runs in by_row.items():
            poet = runs["poet"]
            poets_seen.add(poet)
            m = majority([runs.get(0), runs.get(1)])
            if m is None:
                continue
            slot = per_poet.setdefault(poet, [0, 0])
            slot[0 if m else 1] += 1
        agg[subtype_key] = per_poet

    poet_totals = {}
    for per_poet in agg.values():
        for poet, (tp, fn) in per_poet.items():
            poet_totals[poet] = poet_totals.get(poet, 0) + tp + fn
    poets = sorted(poets_seen, key=lambda p: -poet_totals.get(p, 0))

    def subtype_sort_key(k):
        parent = parent_short(SUBTYPE_DEFS[k]["parent"])
        n_total = sum(sum(v) for v in agg[k].values())
        return (PARENT_ORDER.index(parent) if parent in PARENT_ORDER else 99, -n_total)

    subtypes_ordered = sorted(agg.keys(), key=subtype_sort_key)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "by_poet_subtype_fewshot"
    ws.sheet_view.rightToLeft = True

    header1 = ["", ""]
    header2 = ["الفرع", "العنصر_الرئيسي"]
    for poet in poets:
        header1 += [poet, "", "", ""]
        header2 += ["n", "TP", "FN", "Recall"]
    header1 += ["الإجمالي", "", "", ""]
    header2 += ["n", "TP", "FN", "Recall"]
    ws.append(header1)
    ws.append(header2)
    for c in ws[1]:
        c.font = Font(bold=True)
    for c in ws[2]:
        c.font = Font(bold=True)

    for subtype_key in subtypes_ordered:
        parent = parent_short(SUBTYPE_DEFS[subtype_key]["parent"])
        per_poet = agg[subtype_key]
        row = [subtype_key, parent]
        total_tp = total_fn = 0
        for poet in poets:
            tp, fn = per_poet.get(poet, [0, 0])
            n = tp + fn
            row += [n or None, tp or None, fn or None, round(tp / n, 3) if n else None]
            total_tp += tp
            total_fn += fn
        n_total = total_tp + total_fn
        row += [n_total, total_tp, total_fn, round(total_tp / n_total, 3) if n_total else None]
        ws.append(row)

    n_cols = 2 + 4 * (len(poets) + 1)
    for i in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12

    col = 3
    for poet in poets + ["الإجمالي"]:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 3)
        col += 4
    ws.freeze_panes = "C3"

    note = wb.create_sheet("ملاحظة")
    note.sheet_view.rightToLeft = True
    note.append(["هذا الجدول من تجربة few-shot (13 فرعًا فقط، الفروع التي توفر لها n>=10 لعزل مثالين إضافيين)."])
    note.append(["عدد كل فرع هنا أقل بـ2 من جدول one-shot المقابل، لأن السطرين المستخدَمين كمثالين إضافيين استُبعدا من الاختبار."])
    note.column_dimensions["A"].width = 110

    out_path = RESULTS_DIR / "results_by_poet_and_subtype_fewshot.xlsx"
    wb.save(out_path)
    print(f"written to {out_path}")
    print(f"{len(subtypes_ordered)} subtypes x {len(poets)} poets")


if __name__ == "__main__":
    main()
