"""Leave-one-out one-shot sensitivity probe: كناية عن صفة, جرير بن عطيّة's
17 ground-truth lines only. For each of the 17 lines, use IT as the single
one-shot demonstration (replacing the textbook example in subtype_defs.py)
and test on the other 16 - 17 configs x 16 test rows x 2 runs = 544 calls.
Measures how much a single one-shot's recall depends on WHICH example was
shown, holding the subtype/poet/prompt-structure fixed.

Usage:
    python run_leave_one_out_jarir.py --dry-run
    python run_leave_one_out_jarir.py --workers 10
"""
import argparse
import json
import re
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from tqdm import tqdm

sys.path.insert(0, str(Path(__file__).resolve().parent))

from phase3_config import BY_SUBTYPE_DIR, MODELS, MODEL_KEY, N_RUNS, TEMPERATURE, RESULTS_DIR, ROOT  # noqa: E402
from run_subtype_check import get_verse_block, parse_verse_nos, _parse_response  # noqa: E402
from providers_fixed import call_google_fixed  # noqa: E402

load_dotenv(ROOT / ".env")

SUBTYPE_KEY = "كناية عن صفة"
PARENT = "كناية (Periphrasis)"
DEFINITION_AR = "كناية عن صفة: يُشار بها إلى صفة أو خُلُق معيّن دون ذكر الصفة مباشرة، بذكر أمر يلازمها عادة."
POET_FILTER = "جرير بن عطيّة"

SYSTEM_TEMPLATE = """You are an Arabic rhetoric analysis engine. Follow instructions strictly.

================================
SUB-ELEMENT DEFINITION (REFERENCE)
================================
Below is EXACTLY ONE rhetorical sub-element (a specific branch of a broader
Arabic rhetorical element). Learn it deeply using ONLY the definition and
example given here. Do NOT use outside knowledge, and do NOT confuse it with
sibling sub-elements under the same parent element. You will apply this
definition to ONE verse given in the user's message.

--------------------------------
SUB-ELEMENT
--------------------------------
Parent element: {parent}
Arabic name: {name}
Definition (Arabic): {definition_ar}
Example (from the annotated corpus - verse: {ex_verse}): {ex_note}
"""

from subtype_prompts import ANALYSIS_INSTRUCTIONS  # noqa: E402


def slugify(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|]', "", name)
    return s.replace(" ", "_")


def load_jarir_rows():
    path = BY_SUBTYPE_DIR / f"{slugify(SUBTYPE_KEY)}.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["rows"]
    header = [c.value for c in ws[1]]
    idx = {n: i for i, n in enumerate(header)}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[idx["الشاعر"]] != POET_FILTER:
            continue
        row = {
            "row_id": r[idx["#"]], "verse_no": r[idx["البيت"]], "poet": r[idx["الشاعر"]],
            "span": r[idx["التعبير البلاغي"]], "explanation": r[idx["الشرح"]],
            "poem_key": r[idx["القصيدة_مفتاح"]],
        }
        verse_nos = parse_verse_nos(row["verse_no"])
        row["verse_text"] = get_verse_block(row["poem_key"], verse_nos)
        rows.append(row)
    return rows


def build_messages(demo_row, test_verse_text):
    note = f"the sub-element appears at «{demo_row['span']}»" if demo_row["span"] else "the sub-element is present"
    if demo_row["explanation"]:
        note += f" — {demo_row['explanation']}"
    system = SYSTEM_TEMPLATE.format(
        parent=PARENT, name=SUBTYPE_KEY, definition_ar=DEFINITION_AR,
        ex_verse=demo_row["verse_text"], ex_note=note,
    )
    user = ANALYSIS_INSTRUCTIONS.format(verse=test_verse_text)
    return system, user


def _do_call(demo_row, test_row, run_idx):
    system, user = build_messages(demo_row, test_row["verse_text"])
    result = call_google_fixed(system, user, TEMPERATURE)
    parsed = _parse_response(result["raw_text"])
    record = {
        "demo_row_id": demo_row["row_id"], "demo_verse_no": demo_row["verse_no"],
        "test_row_id": test_row["row_id"], "test_verse_no": test_row["verse_no"],
        "run_idx": run_idx, **parsed,
    }
    return record, result["usage"].get("input_tokens") or 0, result["usage"].get("output_tokens") or 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--workers", type=int, default=10)
    args = ap.parse_args()

    rows = load_jarir_rows()
    n = len(rows)
    print(f"{n} جرير rows for {SUBTYPE_KEY}")

    tasks = []
    for i, demo_row in enumerate(rows):
        test_rows = [r for j, r in enumerate(rows) if j != i]
        for test_row in test_rows:
            for run_idx in range(N_RUNS):
                tasks.append((i, demo_row, test_row, run_idx))

    print(f"{n} configs x {n-1} test rows x {N_RUNS} runs = {len(tasks)} calls planned")
    if args.dry_run:
        return

    model_spec = MODELS[MODEL_KEY]
    total_in_tok = total_out_tok = 0
    n_done = n_errors = 0
    records = []
    bar = tqdm(total=len(tasks), desc="loo jarir", unit="call")

    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = {
            pool.submit(_do_call, demo_row, test_row, run_idx): i
            for i, demo_row, test_row, run_idx in tasks
        }
        for fut in as_completed(futures):
            i = futures[fut]
            try:
                record, in_tok, out_tok = fut.result()
            except Exception as e:
                n_errors += 1
                print(f"  ! call failed (config {i}): {e}", flush=True)
                bar.update(1)
                continue
            record["demo_config_idx"] = i
            records.append(record)
            total_in_tok += in_tok
            total_out_tok += out_tok
            n_done += 1
            bar.update(1)
            if n_done % 50 == 0 or n_done == len(tasks):
                print(f"PROGRESS {n_done+n_errors}/{len(tasks)} ({100*(n_done+n_errors)/len(tasks):.0f}%) - {n_errors} errors", flush=True)
    bar.close()

    out_path = RESULTS_DIR / "raw" / "loo_jarir_kinaya_sifa.jsonl"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        for rec in sorted(records, key=lambda r: (r["demo_config_idx"], r["test_row_id"], r["run_idx"])):
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    price_in = model_spec["price_in"] / 1_000_000
    price_out = model_spec["price_out"] / 1_000_000
    cost = total_in_tok * price_in + total_out_tok * price_out
    print(f"\ntotal API calls: {n_done} ok, {n_errors} failed")
    print(f"total tokens: in={total_in_tok} out={total_out_tok}")
    print(f"estimated cost ({model_spec['display']}): ${cost:.4f}")
    print(f"raw results written to {out_path}")

    meta_path = RESULTS_DIR / "loo_jarir_meta.json"
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump({"rows": rows, "cost_usd": cost, "total_calls": n_done, "total_errors": n_errors}, f, ensure_ascii=False, indent=1)
    print(f"meta written to {meta_path}")


if __name__ == "__main__":
    main()
