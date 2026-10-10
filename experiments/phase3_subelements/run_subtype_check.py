"""Runner for Phase 3 sub-element yes/no checks.

For each confirmed subtype (SUBTYPE_DEFS), loads its ground-truth rows from
data/ground_truth/by_subtype/<slug>.xlsx, resolves the full verse text by
(poem_key, verse_no), and asks the chosen model (Gemini 3.5 Flash, see
phase3_config.MODEL_KEY) whether THIS specific sub-element is present in THIS
verse - N_RUNS=2 times per row, per the project's 2-run minimum rule.

Usage:
    python run_subtype_check.py --dry-run              # no API calls, just verify verse lookup + counts
    python run_subtype_check.py --split dev             # DEV trial (Imru2 + Nabigha only)
    python run_subtype_check.py --split all             # full run (DEV + TEST)
    python run_subtype_check.py --split dev --subtypes "طباق إيجاب" "كناية عن صفة"
    python run_subtype_check.py --split all --workers 12   # parallel calls (default 10)

All (row, run_idx) API calls across every requested subtype are independent,
so they're dispatched through one shared thread pool (--workers, default 10)
rather than sequentially - this is the only thing that changed for the full
run per the user's request to speed it up after watching the DEV trial run
sequentially.
"""
import argparse
import json
import re
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from tqdm import tqdm

sys.path.insert(0, str(Path(__file__).resolve().parent))

from phase3_config import POEMS, MODELS, MODEL_KEY, N_RUNS, TEMPERATURE, BY_SUBTYPE_DIR, RESULTS_DIR, ROOT  # noqa: E402
from subtype_defs import SUBTYPE_DEFS  # noqa: E402
from subtype_prompts import build_messages  # noqa: E402
from providers import CALLERS  # noqa: E402  (kept for non-Gemini providers, unused for MODEL_KEY="google")
from providers_fixed import call_google_fixed  # noqa: E402

load_dotenv(ROOT / ".env")

FENCE_RE = re.compile(r"^```(?:json)?\s*|\s*```$", re.MULTILINE)

_POEM_LINES_CACHE = {}


def get_verse_text(poem_key: str, verse_no: int) -> str:
    """Raw line as-is (including the 'N) "..."' numbering) - matches Phase 1's
    per_verse_prompts.get_verses() convention exactly; the model has always
    been fed lines in this raw form in this project."""
    if poem_key not in _POEM_LINES_CACHE:
        text = POEMS[poem_key]["poetry_path"].read_text(encoding="utf-8").strip()
        _POEM_LINES_CACHE[poem_key] = [ln for ln in text.split("\n") if ln.strip()]
    lines = _POEM_LINES_CACHE[poem_key]
    idx = verse_no - 1
    if not (0 <= idx < len(lines)):
        raise ValueError(f"verse_no {verse_no} out of range for {poem_key} ({len(lines)} verses)")
    return lines[idx]


def parse_verse_nos(cell) -> list:
    """The ground truth's 'البيت' column is usually a single verse number, but
    ~30 rows (out of ~700) use ranges ('9 - 10') or comma-separated groups
    ('4,10 - 11,13,18') when the rhetorical instance spans multiple verses.
    Returns a sorted list of unique verse numbers."""
    if isinstance(cell, (int, float)):
        return [int(cell)]
    nos = set()
    for token in str(cell).split(","):
        token = token.strip()
        if not token:
            continue
        if "-" in token:
            a, b = token.split("-")
            nos.update(range(int(a.strip()), int(b.strip()) + 1))
        else:
            nos.add(int(token))
    return sorted(nos)


def get_verse_block(poem_key: str, verse_nos: list) -> str:
    return "\n".join(get_verse_text(poem_key, v) for v in verse_nos)


def slugify(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|]', "", name)
    return s.replace(" ", "_")


def load_subtype_rows(subtype_key: str, split: str):
    path = BY_SUBTYPE_DIR / f"{slugify(subtype_key)}.xlsx"
    if not path.exists():
        raise FileNotFoundError(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["rows"]
    header = [c.value for c in ws[1]]
    idx = {name: i for i, name in enumerate(header)}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        dev_test = r[idx["DEV_TEST"]]
        if split == "dev" and dev_test != "DEV":
            continue
        if split == "test" and dev_test != "TEST":
            continue
        rows.append({
            "row_id": r[idx["#"]],
            "verse_no": r[idx["البيت"]],
            "poet": r[idx["الشاعر"]],
            "span": r[idx["التعبير البلاغي"]],
            "explanation": r[idx["الشرح"]],
            "poem_key": r[idx["القصيدة_مفتاح"]],
            "dev_test": dev_test,
        })
    return rows


def _parse_response(raw_text: str) -> dict:
    cleaned = FENCE_RE.sub("", raw_text.strip()).strip()
    try:
        obj = json.loads(cleaned)
        return {
            "exists": obj.get("exists"), "appearance": obj.get("appearance"),
            "reasoning": obj.get("reasoning"), "parse_error": False,
        }
    except json.JSONDecodeError:
        return {"exists": None, "appearance": None, "reasoning": None, "parse_error": True, "raw": cleaned[:300]}


def _do_call(caller, subtype_key, row, verse_text, run_idx):
    system, user = build_messages(subtype_key, verse_text)
    result = caller(system, user, TEMPERATURE)
    parsed = _parse_response(result["raw_text"])
    record = {
        "subtype": subtype_key, "run_idx": run_idx, **row,
        "verse_text": verse_text, **parsed, "model_id": result["model_id"],
    }
    return record, result["usage"].get("input_tokens") or 0, result["usage"].get("output_tokens") or 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--split", choices=["dev", "test", "all"], default="dev",
                     help="'test' = TEST-only (use this for the follow-up run after DEV, to avoid re-spending on DEV rows); 'all' = DEV+TEST together (only for a from-scratch run)")
    ap.add_argument("--dry-run", action="store_true", help="no API calls; verify verse lookup + counts only")
    ap.add_argument("--subtypes", nargs="*", default=None, help="restrict to these subtype keys (default: all 23)")
    ap.add_argument("--workers", type=int, default=10, help="parallel API calls (default 10)")
    args = ap.parse_args()

    subtypes = args.subtypes or list(SUBTYPE_DEFS.keys())
    model_spec = MODELS[MODEL_KEY]
    # CALLERS["google"] undercounts Gemini cost (drops thoughts_token_count) - see
    # providers_fixed.py docstring. Use the fixed caller for google; fall back to
    # the shared Phase 1 callers for any other provider key.
    caller = call_google_fixed if MODEL_KEY == "google" else CALLERS[MODEL_KEY]

    out_dir = RESULTS_DIR / "raw"
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1) resolve every (subtype, row) up front so a lookup error fails fast, before spending
    tasks = []  # (subtype_key, row, verse_text, run_idx)
    per_subtype_n = {}
    for subtype_key in subtypes:
        try:
            rows = load_subtype_rows(subtype_key, args.split)
        except FileNotFoundError:
            print(f"  ! no subtype file for {subtype_key!r}, skipping")
            continue
        if not rows:
            print(f"  - {subtype_key}: 0 rows for split={args.split}, skipping")
            continue
        resolved = []
        for row in rows:
            try:
                verse_nos = parse_verse_nos(row["verse_no"])
                verse_text = get_verse_block(row["poem_key"], verse_nos)
            except Exception as e:
                print(f"  ! verse lookup failed for {subtype_key} row #{row['row_id']} (البيت={row['verse_no']!r}): {e}")
                continue
            resolved.append((row, verse_text))
        per_subtype_n[subtype_key] = len(resolved)
        print(f"=== {subtype_key}: {len(resolved)} rows, split={args.split} ===")
        if args.dry_run:
            for row, verse_text in resolved[:2]:
                print(f"  [example] verse_no={row['verse_no']} poem={row['poem_key']} verse_text={verse_text!r}")
            continue
        for row, verse_text in resolved:
            for run_idx in range(N_RUNS):
                tasks.append((subtype_key, row, verse_text, run_idx))

    if args.dry_run:
        total_planned = sum(per_subtype_n.values()) * N_RUNS
        print(f"\n{sum(per_subtype_n.values())} rows total -> {total_planned} calls planned (N_RUNS={N_RUNS})")
        return

    print(f"\ndispatching {len(tasks)} calls across {args.workers} workers...", flush=True)

    # 2) run all calls in parallel; each result is independent (no shared mutable state
    #    except the counters below, which are only ever touched from this main thread
    #    as futures complete - as_completed() serializes that, so no lock is needed).
    total_in_tok = total_out_tok = 0
    n_done = 0
    n_errors = 0
    records_by_subtype = {k: [] for k in per_subtype_n}

    # progress: a live tqdm bar (nice when run in a real terminal) PLUS an
    # explicit, flushed "PROGRESS n/N" line every ~2% of calls - the bar's
    # carriage-return updates don't survive being captured into a redirected
    # log file, but these plain lines do, so progress stays visible either way.
    progress_step = max(1, len(tasks) // 50)
    bar = tqdm(total=len(tasks), desc=f"phase3 {args.split}", unit="call")

    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = {
            pool.submit(_do_call, caller, subtype_key, row, verse_text, run_idx): subtype_key
            for subtype_key, row, verse_text, run_idx in tasks
        }
        for fut in as_completed(futures):
            subtype_key = futures[fut]
            try:
                record, in_tok, out_tok = fut.result()
            except Exception as e:
                n_errors += 1
                print(f"  ! call failed for {subtype_key}: {e}", flush=True)
                bar.update(1)
                continue
            records_by_subtype[subtype_key].append(record)
            total_in_tok += in_tok
            total_out_tok += out_tok
            n_done += 1
            bar.update(1)
            if n_done % progress_step == 0 or n_done == len(tasks):
                pct = 100 * (n_done + n_errors) / len(tasks)
                print(f"PROGRESS {n_done + n_errors}/{len(tasks)} ({pct:.0f}%) - {n_errors} errors so far", flush=True)
    bar.close()

    # 3) write per-subtype jsonl + build summary, same shape as the sequential path
    summary_rows = []
    for subtype_key, records in records_by_subtype.items():
        if not records:
            continue
        out_path = out_dir / f"{slugify(subtype_key)}_{args.split}.jsonl"
        n_yes = [0] * N_RUNS
        with open(out_path, "w", encoding="utf-8") as f:
            for rec in sorted(records, key=lambda r: (r["row_id"], r["run_idx"])):
                if rec.get("exists") is True:
                    n_yes[rec["run_idx"]] += 1
                f.write(json.dumps(rec, ensure_ascii=False) + "\n")
        n_rows = per_subtype_n[subtype_key]
        summary_rows.append({
            "subtype": subtype_key, "n_rows": n_rows,
            **{f"n_yes_run{i+1}": n_yes[i] for i in range(N_RUNS)},
        })
        print(f"  {subtype_key}: {[f'run{i+1}={n_yes[i]}/{n_rows}' for i in range(N_RUNS)]}")

    price_in = model_spec["price_in"] / 1_000_000
    price_out = model_spec["price_out"] / 1_000_000
    cost = total_in_tok * price_in + total_out_tok * price_out

    print()
    print(f"total API calls: {n_done} ok, {n_errors} failed")
    print(f"total tokens: in={total_in_tok} out={total_out_tok}")
    print(f"estimated cost ({model_spec['display']}): ${cost:.4f}")

    if summary_rows:
        summary_path = RESULTS_DIR / f"summary_{args.split}.json"
        RESULTS_DIR.mkdir(parents=True, exist_ok=True)
        with open(summary_path, "w", encoding="utf-8") as f:
            json.dump({
                "split": args.split, "model": model_spec["display"],
                "total_calls": n_done, "total_errors": n_errors,
                "total_in_tok": total_in_tok, "total_out_tok": total_out_tok,
                "cost_usd": cost, "subtypes": summary_rows,
            }, f, ensure_ascii=False, indent=1)
        print(f"summary written to {summary_path}")


if __name__ == "__main__":
    main()
