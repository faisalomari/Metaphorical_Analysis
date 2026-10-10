"""Slice the Phase 1 ground-truth (data/ground_truth/ali_ana/original/*.xlsx) by
secondary rhetorical element ('العنصر البلاغي الثانوي') into one xlsx per subtype.

Only rows whose main element ('العنصر البلاغي المستعمل') falls inside the four
elements this project studies (طباق, تشبيه, مجاز/استعارة, كناية — see GLOSSARY.md)
are considered. Subtypes with fewer than MIN_ROWS rows are dropped (too thin to
say anything about). Output goes to data/ground_truth/by_subtype/.
"""
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SRC_DIR = ROOT / "data" / "ground_truth" / "ali_ana" / "original"
OUT_DIR = ROOT / "data" / "ground_truth" / "by_subtype"
MIN_ROWS = 2

OUR_MAINS = {"طباق", "تشبيه", "مجاز", "استعارة", "مقابلة", "مقابلة تضاد", "كناية"}

# file name -> (poem key, split) — matches experiments/pilot_3models/config.py POEMS/DEV_POEMS/TEST_POEMS
POEM_INFO = {
    "امرؤ القيس 1.xlsx": ("Imru1", "TEST"),
    "امرؤ القيس 8.xlsx": ("Imru2", "DEV"),
    "النابغة الذبياني.xlsx": ("Nabigha", "DEV"),
    "جرير.xlsx": ("Jarir", "TEST"),
    "ابو تمام.xlsx": ("Abu_Tammam", "TEST"),
}

# best-effort label of the parent main element per the catalogue built earlier
# in this conversation (pptx + PDF glossary). None = not matched to our catalogue.
PARENT_OF_SUBTYPE = {
    "طباق إيجاب": "الطباق",
    "طباق سلب": "الطباق",
    "طباق في المعنى": "الطباق",
    "مقابلة (تضاد)": "الطباق",
    "مقابلة تضاد بالمعنى": "الطباق",
    "مقابلة نحوية": None,  # جهاز مستقل، غير فرع من الطباق

    "تشبيه مطول": "التشبيه",
    "تشبيه قصير": "التشبيه",
    "تشبيه تمثيلي": "التشبيه",
    "تشبيه مركب": "التشبيه",
    "تشبيه ضمني": "التشبيه",
    "تشبيه الشيء بنفسه": "التشبيه",
    "تشبيه غريب": None,  # لم يُعرَّف في مصدرينا
    "تشبيه مطول غريب": None,
    "تشبيه مركب غريب": None,
    "تشبيه مركب مطول": "التشبيه",

    "استعارة فعل": "المجاز",
    "مجاز عقلي": "المجاز",
    "استعارة تمثيلية": "المجاز",
    "مجاز مرسل": "المجاز",
    "استعارة قديمة": "المجاز",
    "استعارة حديثة": "المجاز",
    "استعارة قديمة مرشحة": "المجاز",
    "استعارة فعل مرشحة": "المجاز",
    "استعارة تمثيلية مبنية على الكناية": "المجاز",

    "كناية عن صفة": "الكناية",
    "كناية عن موصوف": "الكناية",
    "كناية لغوية": "الكناية",
    "كناية عن فعل": "الكناية",
    "كناية نسبة": "الكناية",
    "كناية فنية": None,  # بلا تعريف في مصدرينا
    "كناية بسيطة": None,
}


def slugify(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|]', "", name)
    s = s.replace(" ", "_")
    return s


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    groups = {}  # sub -> list of (row_values, header, source_file, poem_key, split)

    for xlsx_path in sorted(SRC_DIR.glob("*.xlsx")):
        fname = xlsx_path.name
        if fname not in POEM_INFO:
            print(f"  ! skipping unrecognized file: {fname}")
            continue
        poem_key, split = POEM_INFO[fname]
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb["Sheet1"]
        header = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            main_el, sub_el = row[5], row[6]
            if isinstance(main_el, str):
                main_el = main_el.strip()
            if isinstance(sub_el, str):
                sub_el = sub_el.strip()
            if main_el not in OUR_MAINS:
                continue
            if not sub_el:
                continue
            groups.setdefault(sub_el, []).append((row, header, fname, poem_key, split))

    manifest_rows = []
    kept, dropped = 0, 0
    for sub_el, rows in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        n = len(rows)
        if n < MIN_ROWS:
            dropped += 1
            manifest_rows.append([sub_el, PARENT_OF_SUBTYPE.get(sub_el), n, 0, 0, "", "مستبعد (n<2)"])
            continue
        kept += 1
        header = rows[0][1]
        out_header = list(header) + ["ملف_المصدر", "القصيدة_مفتاح", "DEV_TEST"]
        out_path = OUT_DIR / f"{slugify(sub_el)}.xlsx"

        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active
        ws_out.title = "rows"
        ws_out.append(out_header)
        ws_out.sheet_view.rightToLeft = True
        n_dev = n_test = 0
        for row, _, fname, poem_key, split in rows:
            ws_out.append(list(row) + [fname, poem_key, split])
            if split == "DEV":
                n_dev += 1
            else:
                n_test += 1
        for col_idx in range(1, len(out_header) + 1):
            ws_out.column_dimensions[get_column_letter(col_idx)].width = 18
        wb_out.save(out_path)

        parent = PARENT_OF_SUBTYPE.get(sub_el)
        has_def = "نعم" if parent else "لا — يحتاج تعريف"
        manifest_rows.append([sub_el, parent, n, n_dev, n_test, out_path.name, has_def])

    manifest_path = OUT_DIR / "_manifest.xlsx"
    wb_m = openpyxl.Workbook()
    ws_m = wb_m.active
    ws_m.title = "manifest"
    ws_m.append(["العنصر_الفرعي", "العنصر_الرئيسي", "عدد_الاسطر", "DEV", "TEST", "ملف_الإخراج", "له_تعريف_ومثال_بالجدول"])
    ws_m.sheet_view.rightToLeft = True
    for r in manifest_rows:
        ws_m.append(r)
    for col_idx in range(1, 8):
        ws_m.column_dimensions[get_column_letter(col_idx)].width = 22
    wb_m.save(manifest_path)

    print(f"kept {kept} subtype files (n>=2), dropped {dropped} (n<2)")
    print(f"manifest written to {manifest_path}")


if __name__ == "__main__":
    main()
