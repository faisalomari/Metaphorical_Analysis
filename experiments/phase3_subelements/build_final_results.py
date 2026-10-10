"""Builds the final, per-subtype deliverables from the full run (--split all,
with the cost-tracking fix and the reasoning field):

  results/phase3_subelements/xlsx_results/<subtype>.xlsx
      One row per ground-truth line: every original ground-truth column,
      the resolved verse text sent to the model, and BOTH runs' full
      response (exists / appearance / reasoning) side by side, plus a
      2-run-majority TP/FN verdict column.

  results/phase3_subelements/final_summary.xlsx
      One row per subtype: n, TP, FN, recall (2-run majority), matching
      analyze.py's per_verse_majority convention (a 1-1 split -> FN).

Reads only from the *_all.jsonl raw files (the complete, reasoning-enabled,
cost-corrected run) - NOT the earlier *_dev.jsonl/*_test.jsonl files, which
were a smaller, superseded run without the reasoning field.
"""
import json
import re

import openpyxl
from openpyxl.utils import get_column_letter

from phase3_config import BY_SUBTYPE_DIR, RESULTS_DIR
from subtype_defs import SUBTYPE_DEFS

RAW_DIR = RESULTS_DIR / "raw"
OUT_DIR = RESULTS_DIR / "xlsx_results"


def slugify(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|]', "", name)
    return s.replace(" ", "_")


def majority(votes):
    votes = [v for v in votes if v is not None]
    if not votes:
        return None
    return sum(votes) > len(votes) / 2


def load_model_responses(slug: str):
    """{(poem_key, row_id): {run_idx: {exists, appearance, reasoning}}}"""
    by_row = {}
    path = RAW_DIR / f"{slug}_all.jsonl"
    if not path.exists():
        return by_row
    with open(path, encoding="utf-8") as f:
        for line in f:
            rec = json.loads(line)
            key = (rec["poem_key"], rec["row_id"])
            by_row.setdefault(key, {})[rec["run_idx"]] = rec
    return by_row


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    summary_rows = []

    for subtype_key, d in SUBTYPE_DEFS.items():
        slug = slugify(subtype_key)
        gt_path = BY_SUBTYPE_DIR / f"{slug}.xlsx"
        if not gt_path.exists():
            continue
        wb_gt = openpyxl.load_workbook(gt_path, data_only=True)
        ws_gt = wb_gt["rows"]
        gt_header = [c.value for c in ws_gt[1]]
        gt_idx = {name: i for i, name in enumerate(gt_header)}
        responses = load_model_responses(slug)

        wb_out = openpyxl.Workbook()
        ws = wb_out.active
        ws.title = "results"
        ws.sheet_view.rightToLeft = True
        ws.append(gt_header + [
            "verse_text",
            "run1_exists", "run1_appearance", "run1_reasoning",
            "run2_exists", "run2_appearance", "run2_reasoning",
            "majority_verdict",
        ])

        tp = fn = 0
        for gt_row in ws_gt.iter_rows(min_row=2, values_only=True):
            row_id = gt_row[gt_idx["#"]]
            poem_key = gt_row[gt_idx["القصيدة_مفتاح"]]
            resp = responses.get((poem_key, row_id), {})
            r1 = resp.get(0, {})
            r2 = resp.get(1, {})
            m = majority([r1.get("exists"), r2.get("exists")])
            if m is True:
                tp += 1
                verdict = "TP"
            elif m is False:
                fn += 1
                verdict = "FN"
            else:
                verdict = "لا استجابة"
            ws.append(list(gt_row) + [
                r1.get("verse_text") or r2.get("verse_text"),
                r1.get("exists"), r1.get("appearance"), r1.get("reasoning"),
                r2.get("exists"), r2.get("appearance"), r2.get("reasoning"),
                verdict,
            ])

        for i in range(1, len(gt_header) + 9):
            ws.column_dimensions[get_column_letter(i)].width = 22 if i <= len(gt_header) + 1 else 30

        # definition sheet, for context when this file is opened standalone
        ws_def = wb_out.create_sheet("تعريف_الفرع")
        ws_def.sheet_view.rightToLeft = True
        ws_def.append(["الحقل", "القيمة"])
        ws_def.append(["الفرع", subtype_key])
        ws_def.append(["العنصر الرئيسي", d["parent"]])
        ws_def.append(["التعريف", d["definition_ar"]])
        ws_def.append(["المثال", d["example"]])
        ws_def.column_dimensions["A"].width = 18
        ws_def.column_dimensions["B"].width = 100

        out_path = OUT_DIR / f"{slug}.xlsx"
        wb_out.save(out_path)

        n = tp + fn
        summary_rows.append({
            "subtype": subtype_key, "parent": d["parent"], "n": n,
            "TP": tp, "FN": fn, "recall": round(tp / n, 3) if n else None,
            "file": out_path.name,
        })

    summary_rows.sort(key=lambda r: -r["n"])
    wb_s = openpyxl.Workbook()
    ws_s = wb_s.active
    ws_s.title = "final_summary"
    ws_s.sheet_view.rightToLeft = True
    ws_s.append(["العنصر_الفرعي", "العنصر_الرئيسي", "n", "TP", "FN", "Recall (2-run majority)", "ملف_النتائج"])
    for r in summary_rows:
        ws_s.append([r["subtype"], r["parent"], r["n"], r["TP"], r["FN"], r["recall"], r["file"]])
    for i in range(1, 8):
        ws_s.column_dimensions[get_column_letter(i)].width = 26
    summary_path = RESULTS_DIR / "final_summary.xlsx"
    wb_s.save(summary_path)

    print(f"{len(summary_rows)} per-subtype xlsx files written to {OUT_DIR}")
    print(f"final summary written to {summary_path}")
    print(f"total rows scored: {sum(r['n'] for r in summary_rows)}")


if __name__ == "__main__":
    main()
