"""Builds the one-shot vs few-shot comparison table for the 13 eligible
subtypes (run_fewshot_check.ELIGIBLE_SUBTYPES).

Fairness requirement: one-shot's numbers here are recomputed EXCLUDING the
2 rows held out for few-shot's demonstrations - otherwise a recall
difference could just be "few-shot's test set is missing 2 rows" rather
than "few-shot changed the model's answers." Both conditions are scored on
the identical (n-2) set of rows.

Also builds the per-subtype few-shot xlsx_results (same shape as
build_final_results.py's one-shot output) and a fewshot_comparison.xlsx.
"""
import json
import re

import openpyxl
from openpyxl.utils import get_column_letter

from phase3_config import BY_SUBTYPE_DIR, RESULTS_DIR
from subtype_defs import SUBTYPE_DEFS
from run_fewshot_check import ELIGIBLE_SUBTYPES, N_HOLDOUT

RAW_DIR = RESULTS_DIR / "raw"
OUT_DIR = RESULTS_DIR / "xlsx_results_fewshot"


def slugify(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|]', "", name)
    return s.replace(" ", "_")


def majority(votes):
    votes = [v for v in votes if v is not None]
    if not votes:
        return None
    return sum(votes) > len(votes) / 2


def load_responses(path):
    by_row = {}
    if not path.exists():
        return by_row
    with open(path, encoding="utf-8") as f:
        for line in f:
            rec = json.loads(line)
            key = (rec["poem_key"], rec["row_id"])
            by_row.setdefault(key, {})[rec["run_idx"]] = rec
    return by_row


def score(by_row, exclude_keys=frozenset()):
    tp = fn = 0
    for key, runs in by_row.items():
        if key in exclude_keys:
            continue
        m = majority([runs.get(0, {}).get("exists"), runs.get(1, {}).get("exists")])
        if m is True:
            tp += 1
        elif m is False:
            fn += 1
    n = tp + fn
    return n, tp, fn, round(tp / n, 3) if n else None


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    comparison_rows = []

    for subtype_key in ELIGIBLE_SUBTYPES:
        slug = slugify(subtype_key)

        # the 2 held-out (poem_key, row_id) keys, read back from the ground truth
        # file itself (same order run_fewshot_check.py used: first N_HOLDOUT rows)
        gt_path = BY_SUBTYPE_DIR / f"{slug}.xlsx"
        wb_gt = openpyxl.load_workbook(gt_path, data_only=True)
        ws_gt = wb_gt["rows"]
        gt_header = [c.value for c in ws_gt[1]]
        gt_idx = {name: i for i, name in enumerate(gt_header)}
        all_gt_rows = list(ws_gt.iter_rows(min_row=2, values_only=True))
        holdout_keys = frozenset(
            (r[gt_idx["القصيدة_مفتاح"]], r[gt_idx["#"]]) for r in all_gt_rows[:N_HOLDOUT]
        )

        oneshot_by_row = load_responses(RAW_DIR / f"{slug}_all.jsonl")
        fewshot_by_row = load_responses(RAW_DIR / f"{slug}_fewshot.jsonl")

        n_os, tp_os, fn_os, recall_os = score(oneshot_by_row, exclude_keys=holdout_keys)
        n_fs, tp_fs, fn_fs, recall_fs = score(fewshot_by_row)

        comparison_rows.append({
            "subtype": subtype_key, "parent": SUBTYPE_DEFS[subtype_key]["parent"],
            "n_test": n_fs, "n_test_oneshot_check": n_os,
            "oneshot_TP": tp_os, "oneshot_FN": fn_os, "oneshot_recall": recall_os,
            "fewshot_TP": tp_fs, "fewshot_FN": fn_fs, "fewshot_recall": recall_fs,
            "delta": round((recall_fs or 0) - (recall_os or 0), 3),
        })

        # per-subtype fewshot xlsx (same shape as the one-shot xlsx_results/)
        wb_out = openpyxl.Workbook()
        ws = wb_out.active
        ws.title = "results"
        ws.sheet_view.rightToLeft = True
        ws.append(gt_header + [
            "verse_text", "run1_exists", "run1_appearance", "run1_reasoning",
            "run2_exists", "run2_appearance", "run2_reasoning", "majority_verdict", "held_out_as_fewshot_example",
        ])
        for gt_row in all_gt_rows:
            key = (gt_row[gt_idx["القصيدة_مفتاح"]], gt_row[gt_idx["#"]])
            if key in holdout_keys:
                ws.append(list(gt_row) + [None, None, None, None, None, None, None, "مثال قياسي (few-shot)", "نعم"])
                continue
            resp = fewshot_by_row.get(key, {})
            r1, r2 = resp.get(0, {}), resp.get(1, {})
            m = majority([r1.get("exists"), r2.get("exists")])
            verdict = "TP" if m is True else "FN" if m is False else "لا استجابة"
            ws.append(list(gt_row) + [
                r1.get("verse_text") or r2.get("verse_text"),
                r1.get("exists"), r1.get("appearance"), r1.get("reasoning"),
                r2.get("exists"), r2.get("appearance"), r2.get("reasoning"),
                verdict, "لا",
            ])
        for i in range(1, len(gt_header) + 10):
            ws.column_dimensions[get_column_letter(i)].width = 24
        wb_out.save(OUT_DIR / f"{slug}.xlsx")

    comparison_rows.sort(key=lambda r: -r["n_test"])
    wb_c = openpyxl.Workbook()
    ws_c = wb_c.active
    ws_c.title = "oneshot_vs_fewshot"
    ws_c.sheet_view.rightToLeft = True
    ws_c.append([
        "العنصر_الفرعي", "العنصر_الرئيسي", "n (نفس الأسطر لكلا التجربتين)",
        "One-shot TP", "One-shot FN", "One-shot Recall",
        "Few-shot TP", "Few-shot FN", "Few-shot Recall", "الفرق (Few - One)",
    ])
    for r in comparison_rows:
        ws_c.append([
            r["subtype"], r["parent"], r["n_test"],
            r["oneshot_TP"], r["oneshot_FN"], r["oneshot_recall"],
            r["fewshot_TP"], r["fewshot_FN"], r["fewshot_recall"], r["delta"],
        ])
    for i in range(1, 11):
        ws_c.column_dimensions[get_column_letter(i)].width = 24

    note = wb_c.create_sheet("ملاحظة")
    note.sheet_view.rightToLeft = True
    note.append(["كل صف يقارن نفس الأسطر تمامًا (n) تحت التجربتين — استُبعد السطران المستخدَمان كمثالين إضافيين في few-shot من عدّ One-shot أيضًا حتى تكون المقارنة عادلة."])
    note.append(["الأمثلة الإضافية (سطران لكل فرع) هي أول سطرين بترتيب ملف ground truth لذلك الفرع — اختيار حتمي غير منتقى يدويًا."])
    note.column_dimensions["A"].width = 120

    out_path = RESULTS_DIR / "fewshot_comparison.xlsx"
    wb_c.save(out_path)
    print(f"comparison written to {out_path}")
    print(f"{len(comparison_rows)} subtypes, per-subtype fewshot xlsx in {OUT_DIR}")
    for r in comparison_rows:
        print(f"{r['subtype']:30s} n={r['n_test']:3d}  one-shot={r['oneshot_recall']}  few-shot={r['fewshot_recall']}  delta={r['delta']:+.3f}")


if __name__ == "__main__":
    main()
