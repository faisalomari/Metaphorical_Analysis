"""Per-poet breakdown of the one-shot results (the complete 25-subtype,
570-row run - not the few-shot follow-up, which only covers 13 subtypes on
a reduced test set and would skew poet totals unfairly).

For each poet, aggregates TP/FN/Recall (2-run majority) per main element
(طباق/تشبيه/مجاز/كناية) plus a per-poet total. Note: امرؤ القيس has 2 poems
in this corpus (Imru1, Imru2) - both pooled under the one poet name, matching
the ground truth's own 'الشاعر' column. No new API calls - pure aggregation
of already-collected raw/*_all.jsonl.
"""
import json
import re

import openpyxl
from openpyxl.utils import get_column_letter

from phase3_config import RESULTS_DIR
from subtype_defs import SUBTYPE_DEFS

RAW_DIR = RESULTS_DIR / "raw"

PARENT_ORDER = ["الطباق", "التشبيه", "المجاز", "الكناية"]


def slugify(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|]', "", name)
    return s.replace(" ", "_")


def parent_short(parent_full: str) -> str:
    # SUBTYPE_DEFS parents are English-labelled ("كناية (Periphrasis) — ...");
    # normalize to the 4 short Arabic labels used elsewhere in this project.
    if parent_full.startswith("طباق"):
        return "الطباق"
    if parent_full.startswith("تشبيه"):
        return "التشبيه"
    if parent_full.startswith("مجاز"):
        return "المجاز"
    if parent_full.startswith("كناية"):
        return "الكناية"
    return "غير مصنف"


def majority(votes):
    votes = [v for v in votes if v is not None]
    if not votes:
        return None
    return sum(votes) > len(votes) / 2


def main():
    # {(poet, parent): [tp, fn]}
    agg = {}
    for subtype_key, d in SUBTYPE_DEFS.items():
        slug = slugify(subtype_key)
        path = RAW_DIR / f"{slug}_all.jsonl"
        if not path.exists():
            continue
        parent = parent_short(d["parent"])
        by_row = {}
        with open(path, encoding="utf-8") as f:
            for line in f:
                rec = json.loads(line)
                key = (rec["poem_key"], rec["row_id"])
                by_row.setdefault(key, {"poet": rec["poet"]})[rec["run_idx"]] = rec.get("exists")
        for key, runs in by_row.items():
            poet = runs["poet"]
            m = majority([runs.get(0), runs.get(1)])
            if m is None:
                continue
            slot = agg.setdefault((poet, parent), [0, 0])
            slot[0 if m else 1] += 1

    poets = sorted({p for (p, _) in agg}, key=lambda p: -sum(sum(v) for (pp, e), v in agg.items() if pp == p))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "by_poet"
    ws.sheet_view.rightToLeft = True
    ws.append(["الشاعر", "العنصر", "n", "TP", "FN", "Recall"])

    for poet in poets:
        total_tp = total_fn = 0
        for parent in PARENT_ORDER:
            tp, fn = agg.get((poet, parent), [0, 0])
            if tp + fn == 0:
                continue
            n = tp + fn
            ws.append([poet, parent, n, tp, fn, round(tp / n, 3)])
            total_tp += tp
            total_fn += fn
        n_total = total_tp + total_fn
        ws.append([poet, "الإجمالي", n_total, total_tp, total_fn, round(total_tp / n_total, 3) if n_total else None])
        ws.append([])  # spacer row between poets

    for i in range(1, 7):
        ws.column_dimensions[get_column_letter(i)].width = 20

    # bold the total rows for readability
    from openpyxl.styles import Font
    for row in ws.iter_rows(min_row=2):
        if row[1].value == "الإجمالي":
            for c in row:
                c.font = Font(bold=True)

    out_path = RESULTS_DIR / "results_by_poet.xlsx"
    wb.save(out_path)
    print(f"written to {out_path}")

    print(f"\n{'الشاعر':20s} {'العنصر':10s} {'n':>4s} {'TP':>4s} {'FN':>4s} {'recall':>7s}")
    for poet in poets:
        total_tp = total_fn = 0
        for parent in PARENT_ORDER:
            tp, fn = agg.get((poet, parent), [0, 0])
            if tp + fn == 0:
                continue
            n = tp + fn
            print(f"{poet:20s} {parent:10s} {n:4d} {tp:4d} {fn:4d} {tp/n:7.3f}")
            total_tp += tp
            total_fn += fn
        n_total = total_tp + total_fn
        print(f"{poet:20s} {'الإجمالي':10s} {n_total:4d} {total_tp:4d} {total_fn:4d} {total_tp/n_total:7.3f}")


if __name__ == "__main__":
    main()
