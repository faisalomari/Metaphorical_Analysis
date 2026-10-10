"""For each few-shot-eligible subtype: how many of the 2 held-out
demonstration examples came from each poet, next to that poet's recall on
the (separate, non-overlapping) test rows for the same subtype. Long format:
one row per (subtype, poet). Pure aggregation, no new API calls.
"""
import json

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from phase3_config import RESULTS_DIR
from subtype_defs import SUBTYPE_DEFS
from build_by_poet import parent_short, majority, PARENT_ORDER
from run_fewshot_check import ELIGIBLE_SUBTYPES, split_holdout_and_test
from build_fewshot_comparison import slugify

RAW_DIR = RESULTS_DIR / "raw"
ALL_POETS = ["امرؤ القيس", "النابغة الذبياني", "جرير بن عطيّة", "أبو تمّام"]


def subtype_sort_key(k):
    parent = parent_short(SUBTYPE_DEFS[k]["parent"])
    return (PARENT_ORDER.index(parent) if parent in PARENT_ORDER else 99, -1)


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "demo_vs_recall"
    ws.sheet_view.rightToLeft = True
    ws.append(["الفرع", "العنصر_الرئيسي", "الشاعر", "عدد_الأمثلة_التوضيحية_من_هذا_الشاعر",
               "n_اختبار", "TP", "FN", "Recall"])
    for c in ws[1]:
        c.font = Font(bold=True)

    for subtype_key in sorted(ELIGIBLE_SUBTYPES, key=subtype_sort_key):
        parent = parent_short(SUBTYPE_DEFS[subtype_key]["parent"])
        holdout, _ = split_holdout_and_test(subtype_key)
        demo_counts = {}
        for h in holdout:
            demo_counts[h["poet"]] = demo_counts.get(h["poet"], 0) + 1

        slug = slugify(subtype_key)
        path = RAW_DIR / f"{slug}_fewshot.jsonl"
        by_row = {}
        with open(path, encoding="utf-8") as f:
            for line in f:
                rec = json.loads(line)
                key = (rec["poem_key"], rec["row_id"])
                by_row.setdefault(key, {"poet": rec["poet"]})[rec["run_idx"]] = rec.get("exists")
        test_agg = {}
        for key, runs in by_row.items():
            poet = runs["poet"]
            m = majority([runs.get(0), runs.get(1)])
            if m is None:
                continue
            slot = test_agg.setdefault(poet, [0, 0])
            slot[0 if m else 1] += 1

        poets_here = sorted(set(ALL_POETS) & (set(demo_counts) | set(test_agg)),
                             key=lambda p: -(sum(test_agg.get(p, [0, 0]))))
        for poet in poets_here:
            n_demo = demo_counts.get(poet, 0)
            tp, fn = test_agg.get(poet, [0, 0])
            n_test = tp + fn
            ws.append([subtype_key, parent, poet, n_demo or None,
                       n_test or None, tp or None, fn or None,
                       round(tp / n_test, 3) if n_test else None])

    for i in range(1, 9):
        ws.column_dimensions[get_column_letter(i)].width = 22
    ws.column_dimensions["A"].width = 34
    ws.freeze_panes = "A2"

    note = wb.create_sheet("ملاحظة")
    note.sheet_view.rightToLeft = True
    note.append(["عدد_الأمثلة_التوضيحية_من_هذا_الشاعر: كم من المثالين المستخدَمين لتوضيح الفرع (few-shot) جاء من قصائد هذا الشاعر (0، 1، أو 2)."])
    note.append(["n_اختبار وما بعده: أداء النموذج على أسطر هذا الشاعر المتبقية (غير المستخدَمة كأمثلة) لهذا الفرع، بعد إضافة تلك الأمثلة للبرومبت."])
    note.append(["الأسطر التوضيحية والأسطر المُختبرة منفصلتان تمامًا لكل شاعر - لا تداخل بينهما."])
    note.column_dimensions["A"].width = 120

    out_path = RESULTS_DIR / "fewshot_demo_examples_vs_recall_by_poet.xlsx"
    wb.save(out_path)
    print(f"written to {out_path}")


if __name__ == "__main__":
    main()
