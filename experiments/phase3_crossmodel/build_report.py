"""Builds the Phase 3 cross-model per-subtype table (Gemini vs GPT-5.4-mini,
naive/strict/token recall) + writes a raw JSON snapshot for reproducibility.

No precision/F1 anywhere in this output - undefined by the experiment's
positive-only-sampling design (see score_crossmodel.py's module docstring).
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from crossmodel_config import PHASE3_RAW_DIR, RAW_DIR, RESULTS_DIR  # noqa: E402
from score_crossmodel import score_all_subtypes, MIN_RELIABLE_N  # noqa: E402
from subtype_defs import SUBTYPE_DEFS  # noqa: E402


def main():
    gemini = score_all_subtypes(PHASE3_RAW_DIR, SUBTYPE_DEFS)
    gpt = score_all_subtypes(RAW_DIR / "openai", SUBTYPE_DEFS)

    rows = []
    for k in SUBTYPE_DEFS:
        if k not in gemini or k not in gpt:
            continue
        g, p = gemini[k], gpt[k]
        rows.append({
            "subtype": k, "parent": SUBTYPE_DEFS[k]["parent"], "n": g["n"],
            "gemini_naive": g["recall_naive"], "gemini_strict": g["recall_strict"], "gemini_token": g["recall_token"],
            "gpt_naive": p["recall_naive"], "gpt_strict": p["recall_strict"], "gpt_token": p["recall_token"],
            "gemini_tp_strict": g["tp_strict"], "gemini_fn_strict": g["fn_strict"],
            "gpt_tp_strict": p["tp_strict"], "gpt_fn_strict": p["fn_strict"],
            "thin": g["n"] < MIN_RELIABLE_N,
        })
    rows.sort(key=lambda r: -r["n"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "per_subtype_recall"
    ws.sheet_view.rightToLeft = True
    header = ["العنصر_الفرعي", "العنصر_الرئيسي", "n (=TP+FN)",
              "Gemini naive", "Gemini strict", "Gemini token",
              "GPT-5.4-mini naive", "GPT-5.4-mini strict", "GPT-5.4-mini token",
              "Gemini TP/FN (strict)", "GPT TP/FN (strict)",
              "Precision", "F1", "ملاحظة"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in rows:
        note = "عيّنة صغيرة جدًا (n<10) - غير موثوقة" if r["thin"] else ""
        ws.append([
            r["subtype"], r["parent"], r["n"],
            r["gemini_naive"], r["gemini_strict"], r["gemini_token"],
            r["gpt_naive"], r["gpt_strict"], r["gpt_token"],
            f"{r['gemini_tp_strict']}/{r['gemini_fn_strict']}", f"{r['gpt_tp_strict']}/{r['gpt_fn_strict']}",
            "غير متاح (لا توجد أمثلة سلبية)", "غير متاح (لا توجد أمثلة سلبية)", note,
        ])
    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20

    note_ws = wb.create_sheet("ملاحظة_منهجية")
    note_ws.sheet_view.rightToLeft = True
    note_ws.append(["كل سطر في هذه التجربة كان مثالًا موجبًا معروفًا مسبقًا لفرعه (Dr. Hussein tagged it as this subtype)."])
    note_ws.append(["لذلك TP و FN و Recall مقاسة بصدق فعليًا، لكل من Gemini و GPT-5.4-mini."])
    note_ws.append(["أما FP و TN و Precision و F1: غير قابلة للحساب من هذه التجربة لأي من الموديلين،"
                     " لأننا لم نسأل أي موديل مطلقًا عن سطر لا يحمل هذا الفرع (لا توجد عيّنة سلبية إطلاقًا)."])
    note_ws.append(["هذا الجدول يضيف أيضًا تصحيحًا منهجيًا: Recall المحسوبة هنا (strict/token) تتحقق من أن"
                     " الامتداد الذي أشار إليه الموديل (appearance) يطابق فعليًا الامتداد المُعلَّم في الحقيقة"
                     " الأرضية (span) - لا تكتفي بإجابة 'نعم' فقط (كما فعل الحساب الأصلي naive/recall_by_subtype.xlsx)."])
    note_ws.append([f"الفروع التي n < {MIN_RELIABLE_N} (عدد صغير جدًا) مُعلَّمة في عمود 'ملاحظة' - أرقامها اتجاهية فقط، غير موثوقة للاستنتاج."])
    note_ws.column_dimensions["A"].width = 110

    out_path = RESULTS_DIR / "per_subtype_table.xlsx"
    wb.save(out_path)
    print(f"written to {out_path}")

    # raw JSON snapshot for reproducibility / the summary.md numbers
    snapshot_path = RESULTS_DIR / "raw" / "crossmodel_scoring_snapshot.json"
    with open(snapshot_path, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=1)
    print(f"written to {snapshot_path}")

    return rows


if __name__ == "__main__":
    main()
